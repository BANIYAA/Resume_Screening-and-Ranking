"""
config_manager.py
─────────────────
Excel Reader / Writer for Dynamic Screening Rules.

Each sheet tab in Screening_Rules.xlsx represents a Job Role.
Columns per sheet: Keyword | Category | Weight | Synonyms
"""

import os
import pandas as pd
from openpyxl import load_workbook

# ──────────────────────────────────────────────
# Constants
# ──────────────────────────────────────────────
RULES_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Screening_Rules.xlsx")

DEFAULT_RULES: dict[str, list[dict]] = {
    "UI_UX": [
        {"Keyword": "figma",          "Category": "Design Tools",    "Weight": 10, "Synonyms": "figma tool, figma design"},
        {"Keyword": "wireframing",    "Category": "Design Skills",   "Weight": 8,  "Synonyms": "wireframe, wire-framing, wireframes"},
        {"Keyword": "user research",  "Category": "Research",        "Weight": 8,  "Synonyms": "ux research, usability research, user study"},
        {"Keyword": "adobe xd",       "Category": "Design Tools",    "Weight": 5,  "Synonyms": "xd, adobexd"},
        {"Keyword": "prototyping",    "Category": "Design Skills",   "Weight": 7,  "Synonyms": "prototype, interactive prototype, rapid prototyping"},
        {"Keyword": "sketch",         "Category": "Design Tools",    "Weight": 5,  "Synonyms": "sketch app, sketchapp"},
        {"Keyword": "usability testing", "Category": "Research",     "Weight": 7,  "Synonyms": "usability test, user testing"},
        {"Keyword": "information architecture", "Category": "Design Skills", "Weight": 6, "Synonyms": "ia, info architecture"},
        {"Keyword": "interaction design", "Category": "Design Skills", "Weight": 7, "Synonyms": "ixd, interaction designer"},
        {"Keyword": "design systems", "Category": "Design Skills",   "Weight": 6,  "Synonyms": "design system, component library"},
    ],
    "Python_Developer": [
        {"Keyword": "python",        "Category": "Programming",     "Weight": 10, "Synonyms": "python3, python 3, cpython"},
        {"Keyword": "django",        "Category": "Frameworks",      "Weight": 8,  "Synonyms": "django rest, drf, django framework"},
        {"Keyword": "flask",         "Category": "Frameworks",      "Weight": 7,  "Synonyms": "flask api, flask framework"},
        {"Keyword": "fastapi",       "Category": "Frameworks",      "Weight": 8,  "Synonyms": "fast api, fast-api"},
        {"Keyword": "pandas",        "Category": "Data Libraries",  "Weight": 7,  "Synonyms": "pandas library, pd"},
        {"Keyword": "numpy",         "Category": "Data Libraries",  "Weight": 6,  "Synonyms": "np, numpy library"},
        {"Keyword": "sql",           "Category": "Databases",       "Weight": 7,  "Synonyms": "mysql, postgresql, postgres, sqlite"},
        {"Keyword": "docker",        "Category": "DevOps",          "Weight": 6,  "Synonyms": "docker container, dockerfile, docker-compose"},
        {"Keyword": "git",           "Category": "Version Control", "Weight": 5,  "Synonyms": "github, gitlab, git version control"},
        {"Keyword": "rest api",      "Category": "Architecture",    "Weight": 6,  "Synonyms": "restful, rest apis, restful api"},
    ],
    "Data_Analyst": [
        {"Keyword": "excel",         "Category": "Tools",           "Weight": 8,  "Synonyms": "ms excel, microsoft excel, spreadsheet"},
        {"Keyword": "tableau",       "Category": "Visualization",   "Weight": 9,  "Synonyms": "tableau desktop, tableau public"},
        {"Keyword": "power bi",      "Category": "Visualization",   "Weight": 9,  "Synonyms": "powerbi, power-bi, microsoft power bi"},
        {"Keyword": "sql",           "Category": "Databases",       "Weight": 9,  "Synonyms": "mysql, postgresql, postgres, sqlite, t-sql"},
        {"Keyword": "python",        "Category": "Programming",     "Weight": 7,  "Synonyms": "python3, python 3"},
        {"Keyword": "statistics",    "Category": "Analytics",       "Weight": 7,  "Synonyms": "statistical analysis, stat, descriptive statistics"},
        {"Keyword": "data visualization", "Category": "Visualization", "Weight": 7, "Synonyms": "data viz, dataviz, charts, dashboarding"},
        {"Keyword": "r programming", "Category": "Programming",     "Weight": 6,  "Synonyms": "r language, rstudio, r studio"},
    ],
}


# ──────────────────────────────────────────────
# Bootstrap: create the Excel file if missing
# ──────────────────────────────────────────────
def ensure_rules_file_exists(filepath: str = RULES_FILE) -> None:
    """Create Screening_Rules.xlsx with default sheets if it doesn't exist."""
    if os.path.exists(filepath):
        return
    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        for role_name, rows in DEFAULT_RULES.items():
            df = pd.DataFrame(rows, columns=["Keyword", "Category", "Weight", "Synonyms"])
            df.to_excel(writer, sheet_name=role_name, index=False)


# ──────────────────────────────────────────────
# READ: list available roles
# ──────────────────────────────────────────────
def get_available_roles(filepath: str = RULES_FILE) -> list[str]:
    """Return the sheet names (job roles) present in the Excel workbook."""
    ensure_rules_file_exists(filepath)
    try:
        xls = pd.ExcelFile(filepath, engine="openpyxl")
        return xls.sheet_names
    except Exception as exc:
        print(f"[config_manager] Error reading roles: {exc}")
        return list(DEFAULT_RULES.keys())


# ──────────────────────────────────────────────
# READ: load rules for a specific role
# ──────────────────────────────────────────────
def load_role_rules(role_name: str, filepath: str = RULES_FILE) -> pd.DataFrame:
    """Load the rules DataFrame for the given role sheet."""
    ensure_rules_file_exists(filepath)
    try:
        df = pd.read_excel(filepath, sheet_name=role_name, engine="openpyxl")
        for col in ("Keyword", "Category", "Weight", "Synonyms"):
            if col not in df.columns:
                df[col] = "" if col != "Weight" else 0
        df["Keyword"] = df["Keyword"].astype(str).str.strip().str.lower()
        df["Weight"] = pd.to_numeric(df["Weight"], errors="coerce").fillna(0).astype(int)
        df["Synonyms"] = df["Synonyms"].astype(str).fillna("")
        return df
    except Exception as exc:
        print(f"[config_manager] Error loading rules for '{role_name}': {exc}")
        return pd.DataFrame(columns=["Keyword", "Category", "Weight", "Synonyms"])


# ──────────────────────────────────────────────
# DICT GENERATION: keyword→weight map + spaCy patterns
# ──────────────────────────────────────────────
def generate_keyword_config(role_name: str, filepath: str = RULES_FILE) -> dict:
    """
    Return a dict with:
      - keyword_weights : {variation_string: weight}
      - entity_patterns : list of dicts for spaCy EntityRuler
      - max_possible_score : sum of all unique keyword weights
    """
    df = load_role_rules(role_name, filepath)

    keyword_weights: dict[str, int] = {}
    entity_patterns: list[dict] = []
    max_possible_score: int = 0

    for _, row in df.iterrows():
        primary = str(row["Keyword"]).strip().lower()
        weight = int(row["Weight"])
        category = str(row.get("Category", "SKILL")).strip()
        synonyms_raw = str(row.get("Synonyms", ""))

        if not primary or primary == "nan":
            continue

        # Collect all variations (primary + synonyms)
        variations = [primary]
        if synonyms_raw and synonyms_raw != "nan":
            for syn in synonyms_raw.split(","):
                syn = syn.strip().lower()
                if syn:
                    variations.append(syn)

        # Map every variation to its keyword weight
        for var in variations:
            keyword_weights[var] = weight

        # Build spaCy EntityRuler patterns (multi-token aware)
        label = f"SKILL_{category.upper().replace(' ', '_')}"
        for var in variations:
            tokens = var.split()
            if len(tokens) == 1:
                entity_patterns.append({"label": label, "pattern": var})
            else:
                entity_patterns.append({
                    "label": label,
                    "pattern": [{"LOWER": t} for t in tokens],
                })

        max_possible_score += weight

    return {
        "keyword_weights": keyword_weights,
        "entity_patterns": entity_patterns,
        "max_possible_score": max_possible_score,
    }


# ──────────────────────────────────────────────
# WRITE: add a new role sheet
# ──────────────────────────────────────────────
def add_new_role(role_name: str, filepath: str = RULES_FILE) -> tuple[bool, str]:
    """
    Create a new sheet in the workbook for the given role.
    Returns (success: bool, message: str).
    """
    ensure_rules_file_exists(filepath)

    # Sanitise the name for Excel sheet compatibility
    safe_name = role_name.strip().replace(" ", "_")[:31]
    if not safe_name:
        return False, "Role name cannot be empty."

    try:
        wb = load_workbook(filepath)
        if safe_name in wb.sheetnames:
            wb.close()
            return False, f"Role '{safe_name}' already exists."

        ws = wb.create_sheet(title=safe_name)
        headers = ["Keyword", "Category", "Weight", "Synonyms"]
        ws.append(headers)

        # Add one placeholder row so structure is clear
        ws.append(["example_skill", "General", 5, "synonym1, synonym2"])

        wb.save(filepath)
        wb.close()
        return True, f"Role '{safe_name}' created successfully."
    except Exception as exc:
        return False, f"Error creating role: {exc}"


# ──────────────────────────────────────────────
# Quick self-test
# ──────────────────────────────────────────────
if __name__ == "__main__":
    ensure_rules_file_exists()
    roles = get_available_roles()
    print(f"Available roles: {roles}")
    for r in roles:
        cfg = generate_keyword_config(r)
        print(f"\n── {r} ──")
        print(f"  Keywords : {len(cfg['keyword_weights'])}")
        print(f"  Patterns : {len(cfg['entity_patterns'])}")
        print(f"  Max Score: {cfg['max_possible_score']}")
